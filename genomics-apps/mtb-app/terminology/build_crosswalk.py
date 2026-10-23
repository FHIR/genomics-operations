"""Build public/data/MTB_TumorTypeCrosswalk.csv from the VSAC value-set spreadsheets in value-sets/.

Each value set maps to a predictor tumorType (blank when the predictor has no matching type) and a
gene-list "Cancer category" (see mtb-clinical-data-design.md section 5.5). The spreadsheets are used
exactly as downloaded from VSAC; to refresh, download newer versions and rerun:

    python terminology/build_crosswalk.py
"""
import csv
import warnings
from pathlib import Path

import openpyxl

# VSAC workbooks have no default style; openpyxl warns about it on every load
warnings.filterwarnings('ignore', message='Workbook contains no default style', module='openpyxl')

HERE = Path(__file__).resolve().parent
OUT = HERE.parent / 'public' / 'data' / 'MTB_TumorTypeCrosswalk.csv'

# value-set file -> (predictor tumorType, gene-list Cancer category)
VALUE_SETS = {
    'BreastCancer.xlsx': ('Breast Cancer', 'Breast Cancer'),
    'NonSmallCellLungCancer.xlsx': ('Non-Small Cell Lung Cancer', 'Non-Small Cell Lung Cancer'),
    'HepatocellularCarcinoma.xlsx': ('', 'Hepatocellular Carcinoma'),
}

# VSAC code system names -> FHIR system URIs
CODE_SYSTEMS = {
    'SNOMEDCT': 'http://snomed.info/sct',
}


def value_set_oid(workbook):
    for row in workbook['Value Set Info'].iter_rows(values_only=True):
        if row and row[0] == 'OID':
            return row[1]
    raise ValueError('no OID in Value Set Info sheet')


def expansion(workbook):
    rows = workbook['Expansion List'].iter_rows(values_only=True)
    for row in rows:
        if row and row[0] == 'Code':
            header = list(row)
            break
    else:
        raise ValueError('no header row in Expansion List sheet')
    for row in rows:
        if row and row[0]:
            yield dict(zip(header, row))


def main():
    out_rows = []
    seen = {}
    for filename, (predictor_tumor_type, gene_list_category) in VALUE_SETS.items():
        workbook = openpyxl.load_workbook(HERE / 'value-sets' / filename, read_only=True)
        oid = value_set_oid(workbook)
        for entry in expansion(workbook):
            system = CODE_SYSTEMS[entry['Code System']]
            code = str(entry['Code'])
            if (system, code) in seen:
                raise ValueError(f'{code} is in both {seen[(system, code)]} and {filename}')
            seen[(system, code)] = filename
            out_rows.append({
                'system': system,
                'code': code,
                'display': entry['Description'],
                'predictorTumorType': predictor_tumor_type,
                'geneListCategory': gene_list_category,
                'valueSetOid': oid,
            })

    with open(OUT, 'w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=list(out_rows[0]))
        writer.writeheader()
        writer.writerows(out_rows)
    print(f'{len(out_rows)} codes -> {OUT}')


if __name__ == '__main__':
    main()
